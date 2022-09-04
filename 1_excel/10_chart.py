# 차트 만들기.
from openpyxl.chart import BarChart, Reference, LineChart
from openpyxl import load_workbook
wb = load_workbook("sample.xlsx")
ws = wb.active

# 차트를 만들기 위해 import
# from openpyxl.chart import 차트종류, Reference
# 어떤 데이터를 차트로 만들것인지 value 정의
# bar_value = Reference(ws, min_row=2, max_row=11, min_col=2, max_col=3)
# bar_chart = BarChart()  # 차트 종류 설정, 바, 라인, pie, ...
# bar_chart.add_data(bar_value)  # 차트 데이터 추가.

# ws.add_chart(bar_chart, "E1")  # 차트넣을 위치 정의

# 계열 1, 2 부분 바꾸기.
# B1:C11 까지의 데이터( 과목 이름 포함)
line_value = Reference(ws, min_row=1, max_row=11, min_col=2, max_col=3)
line_chart = LineChart()
line_chart.add_data(line_value, titles_from_data=True)  # 계열부분이 과목이름으로 바뀐다.
line_chart.title = "성적표"  # 제목
line_chart.style = 20  # 미리 정의된 스타일을 적용, 사용자가 개별지정도 가능.
line_chart.y_axis.title = "점수"  # y축의 제목
line_chart.x_axis.title = "번호"  # x축의 제목
ws.add_chart(line_chart, "E1")
# 구글에 openpyxl 가면 차트부분을 눌러 들어가면 거기서 사용자 지정 스타일을 볼수있다.


wb.save("sam_chart.xlsx")
